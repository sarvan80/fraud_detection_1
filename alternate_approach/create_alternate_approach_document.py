from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


FOLDER = Path(r"c:\Users\thiag\Downloads\Schwab-Fraud detection\alternate_approach")
DOCX_PATH = FOLDER / "alternate_approach_security_detection_report.docx"
MD_PATH = FOLDER / "alternate_approach_security_detection_report.md"


REPORT = """# Alternate Approach to Security and Intrusion Detection

## Executive Summary

The alternate approach presented in the PNG materials is a hybrid security-detection framework that combines three complementary mechanisms: signature-based detection for known attacks, anomaly detection for previously unseen behavior, and rule extraction for transparent, operationally maintainable decisions. The approach is designed for environments where raw security events contain many fields, labels are incomplete or expensive to create, and analysts need decisions that can be explained and converted into operational signatures.

The core design begins with network and security event history. Continuous variables are transformed through discretization and binning, categorical or boolean indicators are created through binarization, and the resulting feature space is reduced before clustering. Unsupervised learning identifies normal traffic clusters and separates unusual behavior. Those cluster-derived labels are then passed to a supervised rule-extraction stage. The extracted rules are filtered by precision and recall thresholds, deduplicated semantically, and deployed as a compact rule set for real-time matching. The final output is a decision-intelligence score that can support alerting, prioritization, and investigation.

This architecture is valuable because it does not force every detection problem into a single model family. Known attacks can be handled with precise signatures; unknown or evolving attacks can be identified by deviations from normal profiles; and high-performing machine-learned behavior can be translated into interpretable rules that security teams can review, modify, and extend.

## 1. Scope and Agenda

The source slides organize the solution around five themes:

1. Security and intrusion-detection use cases
2. Modeling philosophy
3. Results and discrepancy analysis
4. Data infrastructure
5. Execution plan

The materials focus on intrusion detection rather than a narrow account-takeover workflow. The examples use event fields such as destination port, device type, source port, protocol name, device name, event name, NAT-port metadata, source and destination IP fields, event count, severity, and timestamps. These fields represent a feature universe typical of security appliances, network monitoring tools, and event-management platforms.

The intended operating context is a security environment where an IDS must distinguish legitimate traffic from malicious traffic entering or leaving the network. The slides explicitly recognize four outcomes:

- True alarms: malicious traffic correctly detected
- Missed detections: malicious traffic that passes undetected
- False alarms: legitimate traffic classified as malicious
- Correctly accepted traffic: legitimate traffic recognized as good

An effective implementation must optimize more than raw accuracy. It must balance detection coverage, analyst workload, false-alarm cost, missed-threat risk, alert latency, and the ability to explain why an event was escalated.

## 2. The Feature Universe and Data Foundation

The feature-universe slide illustrates the heterogeneity of security events. Some fields are numerical or ordinal, such as source and destination ports, magnitude, severity, credibility, event count, and time differences. Other fields are categorical or identifier-like, including protocol name, device name, device type, event name, category, domain, host, identity, and NAT descriptions. Several fields can be missing or defaulted to values such as N/A, zero, false, null, or an all-zero address.

This heterogeneity creates several data-engineering requirements:

- Normalize field names and event schemas across device types.
- Preserve raw event values for forensic review while producing modeling-safe representations.
- Distinguish a true zero from an unavailable, unknown, or not-applicable value.
- Parse timestamps into event time, hour, day-of-week, and elapsed-time features.
- Encode high-cardinality fields carefully so that rare values do not create unstable rules.
- Track the source device, domain, category, and event identity for lineage and investigation.
- Protect against duplicate events, replayed messages, and inconsistent time zones.

The data infrastructure should retain both event-level observations and derived behavioral views. Event-level data supports evidence and rule auditing. Aggregated views support traffic profiling, burst detection, entity behavior, and cluster formation. A practical storage layout is a raw immutable event layer, a normalized security-event layer, a feature layer, and a decision/outcome layer.

## 3. Feature Engineering and Pseudo-Class Construction

The materials identify a set of features such as destination port, device type, source port, protocol name, device name, mPC event, destination and source NAT fields, source IP, destination IP, time difference, event count, and event name. These attributes can be transformed into a pseudo-class variable when reliable human labels are not available.

The pseudo-class examples shown in the slides include:

- Admin login failure
- Brute-force login
- General authentication failure
- Host login failure
- Information event
- Miscellaneous exploit
- Password-change failure
- SSH login failure
- System-security access removed
- User login failure

These labels are not necessarily final ground truth. They are useful operational categories and candidate behavioral classes. A mature system should distinguish at least three label sources:

1. Observed outcome labels from confirmed incidents, analyst dispositions, or response actions.
2. Weak labels derived from event names, severity, system status, or known attack indicators.
3. Unsupervised cluster labels representing normal or unusual behavioral groups.

The distinction matters. If event-name categories are treated as unquestionable truth, the model may merely reproduce existing logging conventions. If cluster labels are treated as confirmed attacks, normal but rare behavior may be over-alerted. The hybrid framework uses these labels as evidence and then validates them through precision, recall, analyst review, and downstream outcomes.

## 4. Modeling Philosophy: Three Complementary Detection Modes

### 4.1 Signature-based detection

The first mode starts with known attacks and historical network logs. A signature-creation process converts known patterns into matching logic. Incoming events are compared against those signatures, producing alerts for recognized attack behavior. Signature detection is fast, familiar to operations teams, and highly precise when the attack pattern is stable.

Its principal limitation is coverage. A signature cannot detect an attack whose structure or values were not anticipated. It may also become brittle when ports, device identifiers, IP addresses, or event sequences change. Signatures therefore remain essential for known threats but should not be the only detection layer.

### 4.2 Anomaly-based detection

The second mode uses training data to build traffic profiles. Network logs are projected into a feature space, optionally reduced in dimensionality, and grouped through clustering or other unsupervised methods. The system identifies normal traffic clusters and scores events that fall outside those profiles as anomalous.

Anomaly detection is useful for new attacks, misconfigurations, compromised hosts, unusual authentication patterns, and deviations from a system's historical baseline. It is especially valuable when confirmed labels are sparse. However, it can generate false alarms when legitimate behavior is seasonal, rare, or affected by operational change. Cluster stability and analyst feedback are therefore important controls.

### 4.3 Rule-extraction framework

The third mode learns rules from data and applies them to incoming events. It sits between opaque statistical scoring and manually written signatures. The extracted rules can represent conjunctions of interpretable conditions, such as a device-type range combined with an event indicator, source-port range, and NAT-port description.

Rule extraction provides a bridge between discovery and operations. It can identify patterns from data, produce a compact rule set, and expose those patterns to subject-matter experts. Experts can accept, reject, edit, or extend the rules. The result is a living detection library rather than a fixed model artifact.

## 5. Hybrid Rule-Based Solution

The hybrid solution shown in the slides treats rule extraction as a supervised task once classification or pseudo-labeling information exists. The learning phase can be unsupervised: dimensionality reduction and clustering discover normal traffic groups and unusual observations. The resulting class information is forwarded to a rule-extraction algorithm that creates conjunctive rules. These rules classify new incoming traffic and can be analyzed by domain experts.

The architecture is intentionally dynamic. Individual algorithms can be changed without redesigning the full system. Classification-based, nearest-neighbor-based, and clustering-based anomaly techniques can be selected according to data quality, latency needs, and operational goals. This makes it possible to begin with a simple baseline and add more sophisticated components as the event history grows.

A recommended flow is:

1. Select representative training data from transaction and network history.
2. Extract and normalize features.
3. Discretize continuous variables and bin ranges where meaningful.
4. Binarize categorical indicators and security flags.
5. Reduce dimensions to remove redundancy and noise.
6. Cluster events and identify normal traffic groups.
7. Assign pseudo-classes or anomaly labels.
8. Extract candidate conjunctive rules.
9. Filter rules using precision, recall, support, and operational cost.
10. Deduplicate semantically similar rules.
11. Validate the rule set with time-separated test data.
12. Deploy the approved rules for real-time matching and decision scoring.

## 6. Classifier-Based Signature Creation

The classifier-based signature slide shows an ensemble approach. Input features X and target values y are passed to a bagging estimator containing decision trees and regression trees. The ensemble produces a set of logical rules. Precision and recall thresholds act as a filter, retaining high-performing rules. Semantic rule deduplication then removes redundant or equivalent rules and produces a set of high-performing, heterogeneous rules.

This approach has several benefits:

- Bagging reduces sensitivity to individual training samples.
- Tree-based models naturally express threshold and conjunction logic.
- Precision and recall thresholds make the business trade-off explicit.
- Deduplication keeps the rule library manageable.
- Heterogeneous rules can cover multiple attack families rather than overfitting to one template.

Rule governance should record each rule's source model, training window, feature conditions, support, precision, recall, first-seen date, owner, approval status, and retirement criteria. A rule should not be promoted solely because it performs well on a static validation set. It should also be evaluated for stability across time, devices, environments, and traffic segments.

## 7. Interpreting the Rule Signatures

The example rule signatures use thresholds on fields such as device name, device type, mPC event, source port, and source-post-NAT-port description. The exact thresholds are less important than the structure: each rule is a conjunction of conditions that jointly identifies a behavior pattern.

For operational use, every generated rule should be translated into a human-readable explanation:

- Which fields triggered the rule?
- What behavior does the conjunction represent?
- How often does it occur in legitimate traffic?
- Which attack or anomaly class is associated with it?
- What evidence supports the alert?
- What action should an analyst or automated system take?

Rules should be tested for sensitivity to missing values, schema changes, numeric boundary effects, and categorical drift. Thresholds should be versioned rather than silently overwritten. When a rule changes, the change should be linked to a validation run and an owner decision.

## 8. Results and Discrepancy Analysis

The discrepancy slide compares unsupervised cluster identifiers with event-classification labels. The table includes categories such as general authentication failed, user login failure, system security access removed, brute-force login, and miscellaneous exploit across clusters 0, 1, 2, and 3.

This comparison is valuable because it reveals where clustering and classification agree or disagree. A category concentrated in one cluster may represent a coherent behavioral pattern. A category spread across several clusters may reflect multiple underlying causes. A cluster containing several event classes may represent a shared operational context rather than a single attack type.

The analysis should go beyond totals:

- Calculate purity of each cluster by event class.
- Calculate entropy to identify mixed clusters.
- Compare cluster assignments across time windows.
- Investigate high-risk classes that appear in otherwise normal clusters.
- Review high-volume classes that generate disproportionate alert load.
- Validate whether clusters correspond to device, protocol, user, or network segments.

Discrepancy is not automatically a failure. It can identify new attack families, inconsistent labeling, overloaded event names, or missing features. The right response is investigation and refinement rather than forcing every cluster to match a pre-existing label.

## 9. Security Use Cases and Operational Outcomes

The security use-case slide emphasizes that intrusion-detection systems are only valuable when they provide elite detection capabilities and accurately distinguish legitimate from malicious traffic. The hybrid approach supports multiple use cases:

- Brute-force and repeated login failure detection
- Suspicious administrative access
- SSH and remote-service abuse
- Credential or authentication anomalies
- Exploit-like event combinations
- Abnormal port and protocol behavior
- Host or device behavior that departs from its baseline
- Unknown or malformed URI/header patterns
- Access removal or account-state changes that indicate compromise

Each use case should define an action policy. Low-confidence anomalies may be logged and enriched. Medium-confidence events may be grouped or rate-limited for analyst review. High-confidence signatures may trigger blocking, isolation, step-up verification, or incident creation. The score should be accompanied by reason codes and supporting event evidence.

## 10. Execution Plan

### Phase 1: Data readiness

Inventory event sources, confirm field semantics, establish time synchronization, identify missing-value conventions, and preserve raw events. Create a common schema for network, device, identity, protocol, and event metadata.

### Phase 2: Baseline profiling

Build normal-traffic profiles by device, protocol, source, destination, user or service, and time window. Measure event frequency, unique-value growth, missingness, and duplicate rates. Establish baseline alert volumes.

### Phase 3: Unsupervised discovery

Apply scaling, binning, binarization, dimensionality reduction, and clustering. Profile clusters and review whether they represent operational segments, normal behavior, or suspicious deviations.

### Phase 4: Rule extraction

Train ensemble classifiers against confirmed or pseudo-labeled classes. Extract candidate rules, tune precision and recall thresholds, remove redundant rules, and produce explanation-ready rule metadata.

### Phase 5: Controlled deployment

Run rules in shadow mode, compare alerts with current controls, measure false alarms and analyst acceptance, and calibrate thresholds. Promote rules gradually with rollback capability.

### Phase 6: Continuous improvement

Add signatures from new incidents, refresh profiles, monitor drift, retire stale rules, and use expert feedback to improve the feature and label layers. Expand the approach to additional domains after proving stability in security and intrusion detection.

## 11. Value Proposition and Next Steps

The source materials identify the main value propositions as high accuracy, no mandatory tagging requirement, customization, robustness, support for subject-matter-expert signature additions, and high interpretability. These strengths come from combining data-driven discovery with human-governed rule operations.

The recommended next steps are:

1. Club the generated signatures with machine-learning scores rather than replacing ML entirely.
2. Expand coverage across additional event domains and attack families.
3. Establish an analyst workflow for rule review, approval, tuning, and retirement.
4. Define evaluation metrics by use case, including precision, recall, false-alarm rate, mean time to detect, and alert reduction.
5. Create a feedback loop from investigations and incident outcomes into labels and training data.
6. Test portability to other security domains only after data semantics and decision policies are documented.

## Conclusion

The alternate approach provides a pragmatic path from raw security logs to explainable, real-time intrusion detection. Signature detection supplies precision for known attacks. Unsupervised profiling supplies discovery for unknown behavior. Rule extraction converts useful patterns into transparent logic that security experts can inspect and maintain. Together, these capabilities form a flexible detection architecture that is adaptable, interpretable, and suitable for incremental production adoption.
"""


def add_markdown_paragraph(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.add_run(text)


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Alternate Approach to Security and Intrusion Detection")
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(31, 55, 86)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Hybrid signatures, anomaly detection, clustering, and explainable rule extraction").italic = True
    doc.add_paragraph("Prepared from the PNG materials in the alternate_approach folder only.").alignment = WD_ALIGN_PARAGRAPH.CENTER

    for block in REPORT.split("\n\n"):
        lines = block.splitlines()
        if not lines:
            continue
        first = lines[0]
        if first.startswith("# "):
            continue
        if first.startswith("## "):
            doc.add_heading(first[3:], level=1)
            if len(lines) > 1:
                add_markdown_paragraph(doc, "\n".join(lines[1:]))
        elif first.startswith("### "):
            doc.add_heading(first[4:], level=2)
            if len(lines) > 1:
                add_markdown_paragraph(doc, "\n".join(lines[1:]))
        elif all(line.startswith("- ") for line in lines):
            for line in lines:
                doc.add_paragraph(line[2:], style="List Bullet")
        elif all(line[:2].isdigit() and ". " in line[:4] for line in lines):
            for line in lines:
                doc.add_paragraph(line.split(". ", 1)[1], style="List Number")
        else:
            add_markdown_paragraph(doc, block)

    doc.save(DOCX_PATH)


MD_PATH.write_text(REPORT, encoding="utf-8")
build_docx()
print(f"Created {MD_PATH}")
print(f"Created {DOCX_PATH}")
